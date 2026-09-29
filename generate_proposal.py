import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# Styles
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    h.runs[0].font.name = 'Arial'
    h.runs[0].font.color.rgb = RGBColor(0, 51, 102)

def add_paragraph(text, align='left', bold=False):
    p = doc.add_paragraph()
    if align == 'center':
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    elif align == 'justify':
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    run = p.add_run(text)
    run.bold = bold
    return p

# PAGE 1: COVER PAGE
doc.add_paragraph("\n\n\n\n\n")
title = doc.add_paragraph()
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run = title.add_run("KERTAS CADANGAN:\nSISTEM PENGIMBAS RESIT ZAKAT PINTAR (AI OCR)\nKE ARAH DIGITALISASI DAN PENGURUSAN DATA YANG KOMPREHENSIF")
run.font.size = Pt(18)
run.bold = True
run.font.color.rgb = RGBColor(0, 51, 102)

doc.add_paragraph("\n\n\n")
sub = doc.add_paragraph()
sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run2 = sub.add_run("Disediakan untuk:\nPengarah\nJabatan Urusan Zakat Wakaf & Baitulmal (JUZWAB)")
run2.font.size = Pt(14)
run2.bold = True

doc.add_paragraph("\n\n")
date = doc.add_paragraph()
date.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run3 = date.add_run("Tarikh Pembentangan: [Tarikh Pembentangan]\nKlasifikasi: SULIT & TERHAD")
run3.font.size = Pt(12)
run3.italic = True

doc.add_page_break()

# PAGE 2: RINGKASAN EKSEKUTIF
add_heading("RINGKASAN EKSEKUTIF", level=1)
add_paragraph("Kertas cadangan ini dibentangkan bertujuan untuk memohon kelulusan dan sokongan daripada pihak Pengurusan Tertinggi Jabatan Urusan Zakat Wakaf & Baitulmal (JUZWAB) bagi membangunkan, memproduksi, dan menggunakan Sistem Pengimbas Resit Zakat Pintar yang berasaskan teknologi Kecerdasan Buatan (AI) Optical Character Recognition (OCR). Sistem ini direka bentuk sebagai penyelesaian 'Back and Forth' berimpak tinggi yang berfungsi secara serampang dua mata: (1) Mengarkibkan resit-resit lama secara pukal (Back), dan (2) Memproses pengeluaran resit baharu secara masa nyata (Forth).", align='justify')
add_paragraph("Dengan kekuatan infrastruktur berbasis web (Web-Based Application), sistem ini berkeupayaan untuk diakses secara lancar menerusi pelayar internet (web browser) pada kedua-dua peranti pintar Android dan iOS tanpa memerlukan sebarang proses muat turun dari Google Play Store mahupun Apple App Store. Ini merendahkan halangan penggunaan (barrier to entry) bagi para amil di lapangan.", align='justify')
add_paragraph("Selain itu, inisiatif digitalisasi ini akan memacu JUZWAB ke arah kelestarian alam sekitar (Go Green) dan penjimatan kos percetakan operasi yang drastik. Jika sebelum ini JUZWAB bergantung kepada penggunaan 3 helaian salinan resit (Salinan JUZWAB, Salinan Amil, dan Salinan Pembayar), sistem ini mampu mengurangkan pergantungan tersebut kepada 1 resit fizikal sahaja. Malah, perancangan jangka masa panjang menyasarkan peralihan 100% kepada E-Resit, sekaligus menamatkan keperluan mencetak resit kertas sepenuhnya.", align='justify')
add_paragraph("Lebih daripada sekadar pendigitalan dokumen, aplikasi ini juga telah merangka pelan penambahbaikan (future planning) yang sangat berharga berdasarkan isu realiti di lapangan, terutamanya kemelut pulangan baki wang syiling (sen-sen) kepada pembayar. Menerusi pengiraan pintar terbina di dalam sistem, sebarang lebihan bayaran zakat akan diproses secara automatik sebagai sumbangan 'Sedekah', menjadikan urusan amil lebih mudah, efisien, dan telus.", align='justify')

doc.add_page_break()

# PAGE 3: 1.0 PENGENALAN & LATAR BELAKANG
add_heading("1.0 PENGENALAN DAN LATAR BELAKANG", level=1)
add_paragraph("Jabatan Urusan Zakat Wakaf & Baitulmal (JUZWAB) memainkan peranan yang sangat kritikal dalam memangkin sosioekonomi umat Islam melalui sistem kutipan dan agihan zakat yang sistematik, cekap, dan berlandaskan Syariat. Selari dengan aspirasi kerajaan ke arah kerangka Ekonomi Digital, setiap agensi kutipan hasil dituntut untuk memperkasa proses kerja harian dengan mengadaptasi teknologi baharu yang berupaya meminimumkan kesilapan manusia (human error), mengoptimumkan kos pengurusan, dan memantapkan integriti data kutipan.", align='justify')
add_paragraph("Secara tradisinya, kutipan zakat fitrah mahupun zakat harta oleh amil-amil yang dilantik di seluruh negara amat bergantung kepada buku resit manual. Buku resit yang dilengkapi dengan helaian kertas karbon (carbon copy) ini memerlukan tenaga kerja fizikal yang intensif. Ia bukan sahaja memakan masa untuk ditulis tangan secara berulang kali, tetapi juga mendedahkan jabatan kepada risiko kerosakan fizikal, kehilangan dokumen, dan kos penyimpanan arkib yang semakin meningkat saban tahun.", align='justify')
add_paragraph("Memandangkan kelantangan arus transformasi digital pada masa kini, adalah menjadi suatu keperluan mendesak bagi JUZWAB untuk mengorak langkah memperkenalkan sebuah sistem yang boleh menjembatani jurang antara proses pengutipan tradisional dengan sistem pangkalan data pusat JUZWAB. Sistem Pengimbas Resit Zakat Pintar ini diperkenalkan sebagai 'Proof of Concept' yang nyata, yang telah terbukti berfungsi dan berpotensi besar untuk diangkat sebagai Standard Operating Procedure (SOP) baharu JUZWAB.", align='justify')

# 2.0 PERNYATAAN MASALAH
add_heading("2.0 PERNYATAAN MASALAH", level=1)
add_paragraph("Walaupun sistem manual yang dipraktikkan sekian lama telah berkhidmat dengan baik, namun evolusi masa dan pertambahan jumlah pembayar zakat telah menyingkap beberapa kelemahan operasi yang serius:", align='justify')
p = doc.add_paragraph(style='List Bullet')
p.add_run("Kos Percetakan dan Isu '3 Salinan' (Carbon Copies): ").bold = True
p.add_run("Penggunaan 3 salinan kertas resit (untuk JUZWAB, Amil, dan Pembayar) melibatkan kos percetakan buku resit yang mencecah ratusan ribu ringgit. Di samping itu, kualiti tulisan tembus karbon kerap kali tidak jelas (illegible) di salinan kedua dan ketiga, menyukarkan kerja-kerja semakan data.")
p = doc.add_paragraph(style='List Bullet')
p.add_run("Kelewatan Pemulangan Buku Resit dan 'Data Backlog': ").bold = True
p.add_run("Terdapat kecenderungan sebilangan amil yang lambat menyerahkan kembali buku resit yang telah habis digunakan ke ibu pejabat. Hal ini menyebabkan proses kemasukan data (data entry) tergendala dan tidak dapat mencerminkan kutipan sebenar pada masa nyata (real-time).")
p = doc.add_paragraph(style='List Bullet')
p.add_run("Kapasiti dan Keselamatan Ruang Arkib: ").bold = True
p.add_run("Buku-buku resit fizikal yang dikumpulkan dari seluruh daerah memerlukan ruang penyimpanan berskala besar dan persekitaran yang terkawal untuk mengelakkan kerosakan berpunca daripada kelembapan, anai-anai, mahupun risiko kebakaran.")
p = doc.add_paragraph(style='List Bullet')
p.add_run("Ralat Manusia (Human Error) dalam Kiraan Tunai: ").bold = True
p.add_run("Semasa waktu puncak, amil kerap kali menghadapi tekanan dalam mengira baki pemulangan wang, khususnya bagi nilai sen-sen. Ini sering menimbulkan perselisihan kira-kira di hujung hari.")

doc.add_page_break()

# PAGE 4: 3.0 OBJEKTIF PROJEK
add_heading("3.0 OBJEKTIF PROJEK", level=1)
add_paragraph("Perlaksanaan Sistem Pengimbas Resit Zakat Pintar ini disandarkan kepada objektif-objektif strategik berikut:", align='justify')
doc.add_paragraph("1. Mendigitalkan proses perekodan resit zakat secara serta-merta sejurus transaksi berlaku melalui penggunaan teknologi pengecaman aksara optik (Optical Character Recognition - OCR).", style='List Number')
doc.add_paragraph("2. Mewujudkan keserasian peranti yang menyeluruh (cross-platform compatibility) dengan memastikan sistem ini boleh diakses dengan lancar pada ekosistem Android dan iOS menerusi pelayar internet standard (Chrome, Safari) tanpa kerumitan pemasangan aplikasi.", style='List Number')
doc.add_paragraph("3. Melaksanakan pendekatan dua hala 'Back and Forth': menterjemah maklumat bertulis dari timbunan resit lama (Arkib) menjadi pangkalan data berstruktur, di samping merekod resit-resit yang baharu dikeluarkan dengan kadar segera.", style='List Number')
doc.add_paragraph("4. Mengurangkan jejak karbon (carbon footprint) JUZWAB melalui pelan fasa pengurangan dari 3 helai resit kepada 1 helai, dengan sasaran mutlak sifar kertas (paperless / 100% E-Resit) di masa hadapan.", style='List Number')
doc.add_paragraph("5. Meningkatkan ketelusan (transparency) dan pemantauan aliran tunai di peringkat ibu pejabat menerusi papan pemuka (dashboard) yang dikemas kini pada masa nyata (real-time).", style='List Number')

# 4.0 KEKUATAN & PENYELESAIAN
add_heading("4.0 PENYELESAIAN YANG DICADANGKAN & KEKUATAN APLIKASI", level=1)
add_paragraph("Berpaksikan kepada keperluan operasi Amil di lapangan, sistem ini telah dibangunkan dengan seni bina berasaskan web (Web-Based App / Progressive Web App) yang menawarkan kekuatan-kekuatan berikut:", align='justify')
add_heading("4.1 Sokongan Peranti Universal (Android & iOS)", level=2)
add_paragraph("Sebilangan besar projek perisian kerajaan berhadapan isu 'adoption rate' yang rendah akibat kesukaran pemasangan aplikasi, versi OS yang tidak serasi, atau kekangan ruang storan telefon. Aplikasi yang dicadangkan ini menyelesaikan dilema tersebut kerana ia diakses terus melalui pautan (URL) dalam pelayar web. Paparannya adalah responsif dan memberikan pengalaman penggunaan (User Experience - UX) seolah-olah menggunakan aplikasi native. Segala kemas kini perisian akan disegerakkan terus secara automatik di bahagian pelayan (server) tanpa amil perlu melakukan muat turun kemas kini.", align='justify')

add_heading("4.2 Pendekatan Serampang Dua Mata ('Back and Forth' Scanning)", level=2)
add_paragraph("Mod 'Forth' (Mod Biasa): Digunakan semasa kutipan aktif. Ia menyaring nama pembayar, No. Kad Pintar, kadar zakat, dan secara berterusan menjana E-Resit dengan Kod QR.", align='justify')
add_paragraph("Mod 'Back' (Quick Scan Arkib Pukal): Dikhaskan untuk operasi pembersihan arkib resit lama. Mod ini dirancang secara khusus untuk memintas (bypass) pengimbasan tulisan tangan yang kabur pada medan nama atau tarikh. Ia memfokuskan kepada ekstraksi kod cetakan statik seperti kod kelompok siri DW/CS dan nombor 6 angka. Pendekatan ini secara drastik menyingkatkan masa pemprosesan dari beberapa minit kepada beberapa saat per resit.", align='justify')

doc.add_page_break()

# PAGE 5
add_heading("4.3 Pengecilan Pergantungan Resit Kertas (Towards Paperless)", level=2)
add_paragraph("Evolusi ini dirangka melalui tiga (3) fasa pelaksanaan yang teratur bagi memastikan amil dan pembayar tidak mengalami kejutan budaya (culture shock):", align='justify')
doc.add_paragraph("Fasa 1: Peralihan dari 3 salinan kepada 1 salinan fizikal. Amil menulis di atas 1 helaian sahaja dan memberikannya kepada pembayar, kemudian Amil mengimbas (scan) resit tersebut menggunakan aplikasi ini untuk dihantar ke database JUZWAB.", style='List Bullet')
doc.add_paragraph("Fasa 2 (Terkini): Ciri 'Generate E-Resit'. Pembayar tidak perlu menunggu resit kertas. Setelah butiran diimbas atau dimasukkan, sistem memaparkan Kod QR E-Resit. Pembayar mengimbas QR tersebut untuk mendapatkan salinan PDF resit rasmi ke dalam telefon bimbit mereka.", style='List Bullet')
doc.add_paragraph("Fasa 3: Pemansuhan percetakan resit fizikal. Transaksi akan bersandar 100% kepada sistem digital, sekali gus menyelesaikan isu ketiadaan stok buku resit atau buku resit hilang.", style='List Bullet')

# 5.0 PERBANDINGAN ALIRAN KERJA (WORKFLOW)
add_heading("5.0 PERBANDINGAN ALIRAN KERJA (WORKFLOW)", level=1)
add_heading("5.1 Aliran Kerja Lama (Manual)", level=2)
p = doc.add_paragraph(style='List Number')
p.add_run("Tulis Tangan: ").bold = True
p.add_run("Amil melengkapkan butiran pembayar ke atas 3 salinan kertas berkarbon (Nama, IC, Jumlah Zakat, Bilangan Tanggungan).")
p = doc.add_paragraph(style='List Number')
p.add_run("Pengasingan Salinan: ").bold = True
p.add_run("Salinan pertama diserahkan kepada pembayar. Salinan kedua disimpan oleh amil, dan salinan ketiga dibiarkan di dalam buku untuk pihak ibu pejabat.")
p = doc.add_paragraph(style='List Number')
p.add_run("Penyerahan Buku: ").bold = True
p.add_run("Selesai musim zakat, amil perlu memandu dan hadir ke pejabat secara fizikal untuk menyerahkan buku resit, bersama wang kutipan (sekiranya tunai).")
p = doc.add_paragraph(style='List Number')
p.add_run("Kemasukan Data: ").bold = True
p.add_run("Kerani di pejabat terpaksa menghabiskan masa berminggu-minggu menaip semula (data entry) beribu-ribu resit ke dalam sistem berdasarkan salinan karbon yang adakalanya mustahil untuk dibaca.")

add_heading("5.2 Aliran Kerja Baharu (Sistem AI OCR & E-Resit)", level=2)
p = doc.add_paragraph(style='List Number')
p.add_run("Pengimbasan (Scan): ").bold = True
p.add_run("Amil membuka aplikasi di pelayar web telefon, memfokuskan kamera pada resit fizikal (Fasa Transisi) atau memasukkan rekod secara maya.")
p = doc.add_paragraph(style='List Number')
p.add_run("AI OCR Membaca: ").bold = True
p.add_run("Teknologi AI membaca butiran di atas dokumen dalam masa beberapa saat dan menyiapkannya dalam borang semakan digital.")
p = doc.add_paragraph(style='List Number')
p.add_run("Pengesahan dan Penjanaan QR: ").bold = True
p.add_run("Amil mengesahkan maklumat, tekan 'Hantar'. Pangkalan data JUZWAB dikemas kini secara automatik. Sebuah Kod QR E-Resit dipaparkan di skrin amil.")
p = doc.add_paragraph(style='List Number')
p.add_run("E-Resit Pembayar: ").bold = True
p.add_run("Pembayar mengimbas Kod QR tersebut menggunakan telefon mereka sendiri dan memuat turun Resit Rasmi (PDF). Selesai.")

doc.add_page_break()

# PAGE 6: IMPLIKASI TEKNIKAL
add_heading("6.0 IMPLIKASI TEKNIKAL DAN KEPERLUAN INFRASTRUKTUR", level=1)
add_paragraph("Bagi memastikan cadangan ini menepati piawaian enterprise tanpa membebankan perbelanjaan awam, infrastruktur yang dibina menitikberatkan aspek kos operasi, kelajuan, dan ketahanan terhadap senario luar jangka (resilience).", align='justify')

add_heading("6.1 Penggunaan Data Internet Amil", level=2)
add_paragraph("Kekhuatiran utama bagi amil di kawasan pedalaman adalah penggunaan kuota data mudah alih. Sistem ini telah mengaplikasikan teknologi Mampatan Imej Client-Side (Client-Side Image Compression). Sebelum mana-mana gambar dihantar ke pelayan (server) untuk diproses, saiz gambar akan dimampatkan dan dihadkan di bawah paras 150KB di peranti pengguna. Ini memastikan kelancaran pemuatan naik walaupun dengan kelajuan internet sekadar 3G atau H+, serta menjimatkan pelan data internet amil secara signifikan.", align='justify')

add_heading("6.2 Threshold dan Spesifikasi Pangkalan Data (Database)", level=2)
add_paragraph("Aplikasi ini dilengkapkan dengan enjin pangkalan data yang amat efisien, bertunjangkan Edge Computing menerusi Cloudflare Workers dan D1 SQLite (Atau PostgreSQL). Sistem ini berupaya mengurus ratusan transaksi (Concurrent Requests) pada satu-satu masa tanpa 'Downtime'. Mengambil kira data setiap resit hanya merangkumi sekitar 2-3 KB maklumat teks, JUZWAB mampu menyimpan rekod puluhan tahun lamanya (jutaan baris data) pada kos pengekalan pelayan (server maintenance) yang amat minimum berbanding solusi Cloud tradisional.", align='justify')

add_heading("6.3 Plan B: Kesediaan Mod Luar Talian (Offline Mode)", level=2)
add_paragraph("Sistem zakat lapangan tidak boleh terhenti hanya disebabkan ketiadaan isyarat telekomunikasi. Oleh itu, sistem ini telah dilengkapi dengan infrastruktur 'Progressive Web App' (PWA) dan 'Service Workers'. Dalam situasi 'Blank Spot' (Tiada Internet):", align='justify')
doc.add_paragraph("1. UI dan aplikasi tetap berfungsi seperti biasa tanpa 'Page Crash'.", style='List Bullet')
doc.add_paragraph("2. Setiap imbasan atau pendaftaran resit baharu akan disimpan dengan selamat secara lokal (IndexedDB) di dalam memori telefon bimbit amil tersebut.", style='List Bullet')
doc.add_paragraph("3. Amil akan diberikan notifikasi peringatan berterusan berwarna jingga tentang baki resit yang tertunggak (Offline Sync Queue).", style='List Bullet')
doc.add_paragraph("4. Sebaik sahaja sambungan internet dipulihkan (contoh: apabila amil pulang ke rumah atau kawasan bandar), sistem akan menyelaraskan (sync) ke semua data tertunggak tersebut ke pangkalan data JUZWAB secara automatik di latar belakang.", style='List Bullet')

doc.add_page_break()

# PAGE 7: FUTURE PLANNING / PENAMBAHBAIKAN
add_heading("7.0 PELAN PENAMBAHBAIKAN MASA HADAPAN (FUTURE PLANNING)", level=1)
add_paragraph("Berdasarkan jangkaan sesi soal jawab (Q&A) dan hasil maklum balas fasa rintis (pilot test), beberapa modul tambahan (Add-ons) telah dirangka dan sedia untuk diintegrasikan pada bila-bila masa JUZWAB bersedia:", align='justify')

add_heading("7.1 Modul Pengiraan Pintar Lebihan Wang (Sedekah/Infaq)", level=2)
add_paragraph("Kerap kali berlaku kesulitan kepada Amil Zakat Fitrah yang perlu memulangkan wang baki dalam denominasi sen yang terlalu kecil (contoh: bayaran RM7.00 untuk kadar zakat RM6.50).", align='justify')
add_paragraph("Penyelesaian yang dicadangkan adalah dengan memasukkan 'Modul Auto-Kira Lebihan' di paparan Semakan. Sekiranya pembayar menyerahkan not RM10.00 bagi kadar RM6.50, Amil hanya perlu memasukkan angka RM10.00, dan sistem akan mengemukakan prompt:", align='justify')
add_paragraph("“Terdapat baki RM3.50. Adakah pembayar bersetuju mewakafkan/menyedekahkan baki ini kepada Baitulmal?”", align='justify')
add_paragraph("Jika ya, E-Resit akan menjana pecahan dua item yang jelas secara serentak:", align='justify')
doc.add_paragraph("Item 1: Zakat Fitrah (RM6.50)", style='List Bullet')
doc.add_paragraph("Item 2: Sedekah/Infaq Am (RM3.50)", style='List Bullet')
doc.add_paragraph("Jumlah Terima: RM10.00", style='List Bullet')
add_paragraph("Ciri inovatif ini bukan sahaja menghilangkan sakit kepala amil mencari wang tukar yang remeh, malah mampu membuka pintu aliran dana tambahan berskala besar untuk JUZWAB hasil daripada derma sen-sen secara terkumpul.", align='justify')

add_heading("7.2 Integrasi Mod E-Payment JUZWAB", level=2)
add_paragraph("Sejajar dengan polisi sifar tunai (Cashless), sistem OCR ini berpotensi dikembangkan menjadi platform (Point-Of-Sale / POS) mudah alih Amil. Aplikasi boleh disambungkan dengan Payment Gateway (seperti FPX atau DuitNow QR). Pembayar tidak perlu membawa tunai; kod QR DuitNow dinamik (dengan jumlah yang spesifik, contohnya RM6.50) akan dijana terus di skrin amil. Pemindahan wang akan terus masuk ke akaun bank JUZWAB, dan sebaik pengesahan bayaran diterima API, E-Resit dikeluarkan secara automatik. Ini mengunci ketirisan dana tunai 100%.", align='justify')

doc.add_page_break()

# PAGE 8: ANALISIS KOS & PENUTUP
add_heading("8.0 ANALISIS KOS-FAEDAH (COST-BENEFIT ANALYSIS)", level=1)
add_paragraph("Perbandingan komprehensif membuktikan bahawa Pelaburan Pulangan (Return on Investment - ROI) ke atas sistem ini adalah sangat menguntungkan di pihak Kerajaan dan Jabatan:", align='justify')

add_heading("8.1 Penjimatan Kewangan (Cost Savings)", level=2)
doc.add_paragraph("Penghapusan Kos Percetakan: Penamatan kontrak percetakan ratusan ribu buku resit berkarbon tiga salinan setiap tahun.", style='List Bullet')
doc.add_paragraph("Pengurangan Kos Logistik: Amil tidak lagi perlu hadir beberapa kali untuk urusan penghantaran/pengambilan buku fizikal.", style='List Bullet')
doc.add_paragraph("Pengoptimuman Sumber Manusia: Tugasan pembantu tadbir atau kerani untuk memasukkan semula (key-in) data resit manual dihapuskan, membolehkan kakitangan ditugaskan ke bahagian yang lebih produktif.", style='List Bullet')

add_heading("8.2 Lonjakan Nilai Tanggungjawab dan Imej", level=2)
doc.add_paragraph("Integriti Tertinggi: Menghapuskan kebarangkalian pemalsuan (fraud), resit hilang, atau wang tunai tidak selari (tally) dengan buku.", style='List Bullet')
doc.add_paragraph("Imej Korporat Cemerlang: JUZWAB akan dilihat sebagai sebuah agensi perintis pentadbiran zakat Islam yang moden, mesra alam, telus, dan ke hadapan dalam penguasaan teknologi setaraf institusi perbankan global.", style='List Bullet')

add_heading("9.0 KESIMPULAN", level=1)
add_paragraph("Pengenalan Sistem Pengimbas Resit Zakat Pintar ini bukanlah sekadar suatu pertukaran alat (tool), sebaliknya sebuah transformasi falsafah operasi kerja JUZWAB secara holistik. Ia sebuah penyelesaian kecil dan pantas dilaksanakan (quick win) namun menjanjikan kelancaran operasi yang tiada tolok bandingnya, kemapanan pengurusan data jangka masa panjang, serta memperteguh keyakinan pembayar zakat awam.", align='justify')
add_paragraph("Oleh yang demikian, sokongan dan mandat penuh daripada pihak Pengurusan Tertinggi amatlah ditagih bagi merealisasikan hasrat memproduksi sistem ini untuk fasa implementasi secara rasmi.", align='justify')

doc.add_paragraph("\n\n\n\n")
sign = doc.add_paragraph("Disediakan oleh,\n\n\n....................................................\n(Nama Pegawai / Pasukan Projek)\nJabatan Urusan Zakat Wakaf & Baitulmal")
sign.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.save("Kertas_Cadangan_Pengimbas_Zakat.docx")
