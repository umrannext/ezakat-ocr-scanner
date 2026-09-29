from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

def add_heading(text, level):
    h = doc.add_heading(text, level)
    return h

def add_paragraph(text, align=None):
    p = doc.add_paragraph(text)
    if align == 'justify':
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    return p

# TITLE
title = doc.add_heading("MANUAL PENGGUNA & SOALAN LAZIM (FAQ)\nSISTEM E-ZAKAT OCR JUZWAB", 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

add_heading("SOALAN LAZIM (FAQ)", level=1)

faqs = [
    {
        "q": "1. Apakah yang perlu saya lakukan jika keluar paparan 'Ralat 1101 (Worker threw exception)' ketika cuba menyimpan resit?",
        "a": "Ralat ini berlaku akibat kesesakan isyarat pelayan sementara (Database Connection Limits). Anda hanya perlu klik 'Refresh' pada pelayar web anda, atau kembali ke halaman utama dan cuba tekan 'Simpan' sekali lagi. Anda juga boleh membiarkan aplikasi beroperasi dalam 'Mod Luar Talian' buat sementara waktu."
    },
    {
        "q": "2. Saya terlupa menanda (tick) kotak 'Jana E-Resit'. Adakah resit tersebut hilang?",
        "a": "Tidak. Sistem sentiasa menjana dan menyimpan E-Resit secara automatik. Anda boleh pergi ke tab 'Sejarah', cari resit tersebut, dan klik butang 'E-Resit' berwarna hijau. Resit digital berserta kod QR anda akan dipaparkan semula dan boleh ditunjuk kepada pembayar."
    },
    {
        "q": "3. Muzakki memaklumkan beliau kehilangan gambar E-Resit yang telah dihantar. Bagaimana saya ingin mencari salinan E-Resit tersebut?",
        "a": "Akses tab 'Sejarah' di aplikasi anda. Gunakan ruang carian untuk menaip nama pembayar atau nombor resit. Setelah dijumpai, klik butang 'E-Resit' dan berikan pautan atau tangkap layar (screenshot) resit tersebut kepada Muzakki."
    },
    {
        "q": "4. Nombor resit yang diimbas (OCR) tidak tepat dengan resit fizikal. Apa perlu saya buat?",
        "a": "Sistem ini memberikan anda kebebasan untuk menyunting. Setelah anda menghantar data, pergi ke tab 'Sejarah', cari rekod yang salah tersebut, klik butang 'Edit' (bersebelahan butang Padam), masukkan nombor yang betul dan tekan 'Simpan Kemaskini'."
    },
    {
        "q": "5. Bolehkah saya memadam rekod jika berlaku kesilapan memasukkan resit yang sama (Duplicate)?",
        "a": "Ya, anda boleh menekan butang 'Padam' (ikon tong sampah merah) pada rekod di halaman Sejarah. Sistem akan meminta anda menaip perkataan 'padam' sebagai pengesahan keselamatan sebelum rekod tersebut dipadam sepenuhnya daripada pangkalan data."
    },
    {
        "q": "6. Telefon saya kehabisan ruang memori kerana terlalu banyak imej resit yang diimbas, apa penyelesaiannya?",
        "a": "Pergi ke tab 'Sejarah', di sebelah atas butang 'Refresh', terdapat butang merah ikon tong sampah (Kosongkan Cache). Butang ini akan memadam semua cache imej di dalam telefon anda tanpa memadam rekod data atau pautan E-Resit di Cloud/Pangkalan Data."
    },
    {
        "q": "7. Kenapa paparan 'ISYARAT TERPUTUS (MOD LUAR TALIAN)' muncul di skrin saya?",
        "a": "Ini bermakna telefon bimbit anda berada di kawasan tiada capaian internet (blank spot) atau isyarat lemah. Jangan risau, teruskan kerja menangkap gambar resit dan sistem akan menyimpannya ke dalam cache (Offline Mode). Sistem akan 'Auto-Sync' secara rahsia sebaik sahaja internet disambungkan semula."
    },
    {
        "q": "8. Bagaimana sistem memproses 'Wang Lebihan' menjadi Sedekah secara automatik?",
        "a": "Ketika menyemak imbasan Zakat Fitrah, terdapat ruangan 'Baki'. Jika pembayar ingin menginfaqkan baki tersebut, tandakan (tick) pada 'Jadikan Baki Sebagai Sedekah'. Sistem akan memecahkan transaksi Zakat (jumlah wajib) dan Sedekah (wang baki) ke dalam dua log yang berbeza bagi memudahkan pihak Ibu Pejabat membuat auditan."
    },
    {
        "q": "9. Sistem tiba-tiba tidak boleh log masuk menggunakan Nombor Kad Pengenalan saya, apakah puncanya?",
        "a": "Sila pastikan Nombor Kad Pengenalan anda sah di dalam daftar kakitangan Amil. Sekiranya amaran masih keluar, ini mungkin kerana pihak Admin menyahaktifkan (Deactivate) atau menggantung akaun anda. Sila hubungi Pejabat JUZWAB."
    },
    {
        "q": "10. Apakah perbezaan resit Zakat Harta dengan Zakat Fitrah di dalam rekod?",
        "a": "Di halaman Sejarah, resit Zakat Harta mempunyai tanda atau lencana berwarna OREN dan ditandakan secara berasingan. Imbasan Zakat Harta selalunya mengesan nombor siri 5-angka dan warna kertas putih (seperti dokumen Borang A atau format Lanskap)."
    },
    {
        "q": "11. Bagaimana pihak Admin mendapat akses kepada data kutipan saya?",
        "a": "Setiap butiran yang anda 'Simpan' akan dipancarkan secara langsung (real-time) kepada Papan Pemuka (Dashboard) Ibu Pejabat. Pihak Pengurusan Eksekutif dapat memantau jumlah kutipan dan memuat turun penyata tersebut dalam format jadual Excel (.CSV)."
    },
    {
        "q": "12. Adakah E-Resit digital (QR Code) diiktiraf sebagai bukti sah pembayaran Zakat?",
        "a": "Ya. Kod QR yang dijana menghubungkan ID transaksi unik yang disulitkan secara langsung di pelayan awan JUZWAB, menjadikannya sangat selamat, tahan selamanya, dan diiktiraf setara dengan resit karbon fizikal lama."
    }
]

for item in faqs:
    add_heading(item["q"], level=2)
    add_paragraph(item["a"], align='justify')

doc.add_page_break()
add_heading("PANDUAN LANGKAH DEMI LANGKAH", level=1)
doc.add_paragraph("PANDUAN AMIL (PENGGUNA)", style='Heading 2')
doc.add_paragraph("Langkah 1: Log Masuk - Buka aplikasi melalui pautan rasmi. Masukkan No. Kad Pengenalan dan Kata Laluan anda.", style='List Number')
doc.add_paragraph("Langkah 2: Proses Imbasan - Pilih jenis Zakat. Buka kamera/muat naik gambar resit. Tunggu sistem OCR mengekstrak nombor, jenis kertas dan harga beras.", style='List Number')
doc.add_paragraph("Langkah 3: Semakan - Pada paparan ulasan (Review), pastikan Nombor Resit dan jumlah tanggungan adalah tepat. Anda boleh membetulkan harga sekiranya tersilap bacaan.", style='List Number')
doc.add_paragraph("Langkah 4: Penyimpanan & Jana E-Resit - Klik 'Simpan'. Tetingkap resit digital (beserta QR) akan terbuka automatik jika anda memilih untuk melihatnya.", style='List Number')
doc.add_paragraph("Langkah 5: Menyunting Rekod - Jika terdapat ralat nombor OCR selepas menyimpan, pergi ke halaman 'Sejarah', cari rekod tersebut, klik 'Edit', betulkan nombornya, dan 'Simpan Kemaskini'.", style='List Number')

doc.add_paragraph("\n\nPANDUAN ADMIN (PENTADBIR)", style='Heading 2')
doc.add_paragraph("Langkah 1: Log Masuk Pentadbir - Gunakan ID Admin. Anda akan dihalakan terus ke Papan Pemuka (Dashboard).", style='List Number')
doc.add_paragraph("Langkah 2: Pemantauan Dashboard - Lihat jumlah kutipan keseluruhan secara langsung (real-time) berdasarkan carta dan senarai kemasukan.", style='List Number')
doc.add_paragraph("Langkah 3: Muat Turun Data .CSV - Klik butang Eksport CSV untuk memuat turun jadual Excel bagi keperluan auditan kewangan JUZWAB mengikut penapisan zon dan Tahun Hijrah.", style='List Number')
doc.add_paragraph("Langkah 4: Pengurusan Tetapan - Anda boleh mendaftarkan amil secara pukal (Bulk Upload), menyelaras kadar harga beras (Fitrah), dan menyegerakkan kadar harga Emas semasa (Harta) di tab 'Tetapan'.", style='List Number')


doc.save("Panduan_Pengguna_Dan_FAQ.docx")
print("Saved Panduan_Pengguna_Dan_FAQ.docx")
